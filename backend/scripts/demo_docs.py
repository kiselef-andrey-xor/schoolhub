"""Генерация демо-документов (PDF/DOCX) для seed-данных.

PDF собирается через reportlab; кириллица обеспечивается TTF-шрифтом
(по умолчанию — Noto Sans из backend/assets/fonts, см. NOTICE.md).
DOCX — через python-docx (это валидный OOXML-пакет, открывается в Word/LibreOffice).

Обе библиотеки опциональны: если их нет, seed просто пропустит файлы.
"""
from __future__ import annotations

import io
import os
from dataclasses import dataclass, field
from datetime import date
from pathlib import Path

FONT_DIR = Path(__file__).resolve().parent.parent / "assets" / "fonts"
FONT_PATH = Path(os.getenv("DEMO_FONT_PATH") or (FONT_DIR / "NotoSans-Regular.ttf"))
FONT_BOLD = FONT_DIR / "NotoSans-Bold.ttf"
FONT_ITALIC = FONT_DIR / "NotoSans-Italic.ttf"


@dataclass
class DemoDoc:
    school: str = "ГБОУ «Школа № 1472»"
    kind: str = "ПРИКАЗ"
    number: str = "112-к"
    issued_on: date = field(default_factory=date.today)
    city: str = "г. Москва"
    title: str = ""
    preamble: str = ""
    items: list[str] = field(default_factory=list)
    controls: list[str] = field(default_factory=list)
    signer: str = "Директор школы"
    signer_name: str = "И. А. Соколова"
    acquainted_note: str = (
        "С приказом ознакомлены: лист ознакомления формируется автоматически "
        "в электронной системе SchoolHub."
    )


def build_pdf(doc: DemoDoc) -> bytes | None:
    try:
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
        from reportlab.lib.units import mm
        from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, HRFlowable
    except ImportError:  # pragma: no cover
        return None

    styles = getSampleStyleSheet()
    base = "Helvetica"
    bold = "Helvetica-Bold"
    italic = "Helvetica-Oblique"
    if FONT_PATH.exists():
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.ttfonts import TTFont

        pdfmetrics.registerFont(TTFont("Demo", str(FONT_PATH)))
        base = "Demo"
        if FONT_BOLD.exists():
            pdfmetrics.registerFont(TTFont("Demo-Bold", str(FONT_BOLD)))
            bold = "Demo-Bold"
        else:
            bold = "Demo"
        if FONT_ITALIC.exists():
            pdfmetrics.registerFont(TTFont("Demo-Italic", str(FONT_ITALIC)))
            italic = "Demo-Italic"

    s_school = ParagraphStyle("school", parent=styles["Normal"], fontName=bold, fontSize=11, alignment=1, leading=14)
    s_kind = ParagraphStyle("kind", parent=styles["Normal"], fontName=bold, fontSize=14, alignment=1, spaceBefore=6)
    s_meta = ParagraphStyle("meta", parent=styles["Normal"], fontName=base, fontSize=10, alignment=1, leading=13)
    s_title = ParagraphStyle("title", parent=styles["Normal"], fontName=bold, fontSize=11.5, alignment=1, leading=15, spaceBefore=8)
    s_body = ParagraphStyle("body", parent=styles["Normal"], fontName=base, fontSize=10.5, leading=15, alignment=4, spaceBefore=6)
    s_item = ParagraphStyle("item", parent=s_body, leftIndent=10, firstLineIndent=0, spaceBefore=3)
    s_small = ParagraphStyle("small", parent=styles["Normal"], fontName=italic, fontSize=9, leading=12, textColor=colors.grey)

    buf = io.BytesIO()
    pdf = SimpleDocTemplate(
        buf, pagesize=A4, leftMargin=22 * mm, rightMargin=18 * mm, topMargin=18 * mm, bottomMargin=18 * mm,
        title=f"{doc.kind} №{doc.number}", author=doc.school,
    )
    story = [
        Paragraph(doc.school, s_school),
        HRFlowable(width="100%", thickness=0.8, color=colors.HexColor("#1f3a68"), spaceBefore=4, spaceAfter=8),
        Paragraph(doc.kind, s_kind),
        Paragraph(f"№ {doc.number} от {doc.issued_on:%d.%m.%Y}", s_meta),
        Paragraph(doc.city, s_meta),
        Paragraph(doc.title, s_title),
        HRFlowable(width="100%", thickness=0.4, color=colors.HexColor("#c9d3e3"), spaceBefore=6, spaceAfter=4),
        Paragraph(doc.preamble, s_body),
        Spacer(1, 4),
    ]
    for i, item in enumerate(doc.items, start=1):
        story.append(Paragraph(f"{i}. {item}", s_item))
    if doc.controls:
        story.append(Spacer(1, 6))
        for c in doc.controls:
            story.append(Paragraph(c, s_item))
    story += [
        Spacer(1, 18),
        Paragraph(f"{doc.signer} &nbsp;&nbsp;&nbsp; {doc.signer_name}", s_body),
        Spacer(1, 20),
        Paragraph(doc.acquainted_note, s_small),
    ]
    pdf.build(story)
    return buf.getvalue()


def build_docx(doc: DemoDoc) -> bytes | None:
    try:
        import docx  # python-docx
    except ImportError:  # pragma: no cover
        return None

    d = docx.Document()
    style = d.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = docx.shared.Pt(11)

    p = d.add_paragraph()
    p.alignment = docx.enum.text.WD_ALIGN_PARAGRAPH.CENTER
    p.add_run(doc.school).bold = True

    p = d.add_paragraph()
    p.alignment = docx.enum.text.WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(doc.kind)
    r.bold = True
    r.font.size = docx.shared.Pt(14)

    p = d.add_paragraph()
    p.alignment = docx.enum.text.WD_ALIGN_PARAGRAPH.CENTER
    p.add_run(f"№ {doc.number} от {doc.issued_on:%d.%m.%Y}, {doc.city}")

    p = d.add_paragraph()
    p.alignment = docx.enum.text.WD_ALIGN_PARAGRAPH.CENTER
    p.add_run(doc.title).bold = True

    d.add_paragraph(doc.preamble)
    for i, item in enumerate(doc.items, start=1):
        d.add_paragraph(f"{i}. {item}")
    for c in doc.controls:
        d.add_paragraph(c)

    d.add_paragraph()
    d.add_paragraph(f"{doc.signer}          {doc.signer_name}")
    note = d.add_paragraph(doc.acquainted_note)
    note.runs[0].italic = True
    note.runs[0].font.size = docx.shared.Pt(9)

    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()
